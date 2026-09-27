"""LeadEngine FastAPI application.

Implements ``SPECS/API_CONTRACT.md`` v1 exactly. All product routes live
under ``/api/v1``. Auth is Bearer JWT (standard base64url JSON, HS256,
see ``backend/auth_service.py``).

Conventions (per contract):
- Errors: HTTP status + ``{"error": {"code", "message", "details"}}``.
  Never 200-with-error, never stack traces to clients.
- Pagination: cursor-based, opaque base64url cursor over (sort key, id).
  No OFFSET, no exact totals on list endpoints.
- Sorting: ``sort`` allowlisted per endpoint, ``dir=asc|desc``.
- All SQL parameterized; sort columns and other identifiers allowlisted.
- ``verification_status`` is always ``"unverified"`` until a real
  verification pipeline exists.
- ``lead_score`` is a fit score over unverified data, never verified intent.

Saved searches and favorites are single-tenant local state (no login
required); they live in the ``saved_search`` / ``favorite`` DuckDB tables.
"""

from __future__ import annotations

import base64
import csv
import io
import json
import time
import uuid
from contextlib import asynccontextmanager
from typing import Any, Optional

from fastapi import Depends, FastAPI, Query, Request
from fastapi.exceptions import RequestValidationError
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
from pydantic import BaseModel, Field, field_validator

from backend import auth_service, config, database

API_VERSION = "1.0.0"

# ── Errors ──────────────────────────────────────────────────────────────

class APIError(Exception):
    """Application error rendered as {"error": {"code","message","details"}}."""

    def __init__(self, status: int, code: str, message: str, details: Optional[dict] = None):
        super().__init__(message)
        self.status = status
        self.code = code
        self.message = message
        self.details = details or {}


def _error_body(code: str, message: str, details: Optional[dict] = None) -> dict:
    return {"error": {"code": code, "message": message, "details": details or {}}}


# ── App & lifespan ──────────────────────────────────────────────────────

@asynccontextmanager
async def lifespan(app: FastAPI):
    database.ensure_aux_tables()
    yield


app = FastAPI(title="LeadEngine API", version=API_VERSION, lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=config.CORS_ORIGINS,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.exception_handler(APIError)
async def _api_error_handler(request: Request, exc: APIError):
    return JSONResponse(
        status_code=exc.status,
        content=_error_body(exc.code, exc.message, exc.details),
    )


@app.exception_handler(RequestValidationError)
async def _validation_error_handler(request: Request, exc: RequestValidationError):
    errors = [
        {"loc": list(e.get("loc", [])), "msg": e.get("msg"), "type": e.get("type")}
        for e in exc.errors()
    ]
    return JSONResponse(
        status_code=422,
        content=_error_body("validation_error", "request validation failed",
                            {"errors": errors}),
    )


# ── Pydantic models ─────────────────────────────────────────────────────

class SignupIn(BaseModel):
    email: str = Field(min_length=3, max_length=320)
    password: str = Field(min_length=8, max_length=256)
    name: str = Field(min_length=1, max_length=200)

    @field_validator("email")
    @classmethod
    def _email_valid(cls, v: str) -> str:
        v = v.strip().lower()
        if "@" not in v or "." not in v.split("@")[-1]:
            raise ValueError("email is not valid")
        return v

    @field_validator("name")
    @classmethod
    def _name_stripped(cls, v: str) -> str:
        v = v.strip()
        if not v:
            raise ValueError("name is required")
        return v


class LoginIn(BaseModel):
    email: str = Field(min_length=1, max_length=320)
    password: str = Field(min_length=1, max_length=256)


class UserOut(BaseModel):
    id: str
    email: str
    name: str
    org_id: Optional[str] = None
    role: str
    created_at: float


class TokenOut(BaseModel):
    token: str
    user: UserOut


class CompanySummary(BaseModel):
    id: str
    name: Optional[str] = None
    domain: Optional[str] = None
    website: Optional[str] = None
    industry: Optional[str] = None
    hq_location: Optional[str] = None
    employee_enum: Optional[str] = None
    employees_min: Optional[int] = None
    employees_max: Optional[int] = None
    founded_year: Optional[int] = None
    operating_status: Optional[str] = None
    has_email: bool = False
    has_phone: bool = False
    has_linkedin: bool = False
    quality_score: Optional[float] = None
    lead_score: Optional[float] = None
    verification_status: str = "unverified"


class CompanyDetail(CompanySummary):
    description: Optional[str] = None
    email: Optional[str] = None
    email_valid: Optional[bool] = None
    phone: Optional[str] = None
    facebook: Optional[str] = None
    linkedin: Optional[str] = None
    twitter: Optional[str] = None
    industries: list[str] = Field(default_factory=list)
    locations_raw: Optional[str] = None
    location_tail: Optional[str] = None
    funding_total_usd: Optional[float] = None
    funding_rounds: Optional[int] = None
    last_funding_type: Optional[str] = None
    last_funding_at: Optional[str] = None
    founded_on: Optional[str] = None
    company_type: Optional[str] = None
    ipo_status: Optional[str] = None
    crunchbase_url: Optional[str] = None
    score_signals: list[dict] = Field(default_factory=list)
    source: Optional[str] = None
    observed_at: Optional[str] = None


class CompanyPage(BaseModel):
    items: list[CompanySummary]
    next_cursor: Optional[str] = None
    limit: int


class FacetItem(BaseModel):
    value: str
    count: int


class EmployeeRangeFacet(BaseModel):
    value: str
    label: str
    count: int


class FacetsOut(BaseModel):
    industries: list[FacetItem]
    locations: list[FacetItem]
    operating_statuses: list[FacetItem]
    employee_ranges: list[EmployeeRangeFacet]


class SavedSearchIn(BaseModel):
    name: str = Field(min_length=1, max_length=200)
    filters: dict = Field(default_factory=dict)


class SavedSearchOut(BaseModel):
    id: str
    name: str
    filters: dict
    created_at: float


class SavedSearchList(BaseModel):
    items: list[SavedSearchOut]


class BulkExportIn(BaseModel):
    ids: list[str] = Field(min_length=1, max_length=5000)
    format: str = Field(pattern="^(csv|json|xlsx)$")
    columns: Optional[list[str]] = None

    @field_validator("columns")
    @classmethod
    def _columns_allowlisted(cls, v: Optional[list[str]]) -> Optional[list[str]]:
        if v is None:
            return v
        bad = [c for c in v if c not in EXPORT_COLUMNS]
        if bad:
            raise ValueError(f"columns not allowlisted: {bad}")
        return v


class FavoriteOut(BaseModel):
    company_id: str
    created_at: float


# ── Auth helpers ────────────────────────────────────────────────────────

def _public_user(user: auth_service.User) -> dict:
    return {
        "id": user.id,
        "email": user.email,
        "name": user.name,
        "org_id": user.org_id,
        "role": user.role,
        "created_at": user.created_at,
    }


def get_current_user(request: Request) -> dict:
    """Bearer-token dependency. Raises 401 invalid_token on any problem."""
    auth = request.headers.get("Authorization", "")
    if not auth.startswith("Bearer "):
        raise APIError(401, "invalid_token",
                       "missing or malformed Authorization header (expected 'Bearer <token>')")
    token = auth[7:].strip()
    try:
        payload = auth_service.verify_token(token)
    except ValueError:
        raise APIError(401, "invalid_token", "invalid or expired token")
    user = auth_service.get_user(payload.get("sub", ""))
    if user is None:
        raise APIError(401, "invalid_token", "invalid or expired token")
    return _public_user(user)


# ── Company search core ─────────────────────────────────────────────────

# sort key -> SQL expression (total: NULLs coalesced to a bottom sentinel)
_SORT_EXPRS = {
    "name": "COALESCE(c.name, '')",
    "quality_score": "COALESCE(c.quality_score, -1)",
    "employees_max": "COALESCE(c.employees_max, -1)",
    "founded_year": "COALESCE(c.founded_year, -1)",
}
_DEFAULT_DIRS = {"name": "asc"}  # everything else defaults to desc

SUMMARY_SELECT = (
    "c.id, c.name, c.domain, c.website, c.industry, c.hq_location, "
    "c.employee_enum, c.employees_min, c.employees_max, c.founded_year, "
    "c.operating_status, c.quality_score, c.lead_score, "
    "'unverified' AS verification_status, "
    "(c.email_valid IS TRUE) AS has_email, "
    "(c.phone_normalized IS NOT NULL AND c.phone_normalized <> '') AS has_phone, "
    "(c.linkedin IS NOT NULL AND c.linkedin <> '') AS has_linkedin"
)
_SUMMARY_KEYS = [
    "id", "name", "domain", "website", "industry", "hq_location",
    "employee_enum", "employees_min", "employees_max", "founded_year",
    "operating_status", "quality_score", "lead_score",
    "verification_status", "has_email", "has_phone", "has_linkedin",
]

# Columns a client may request in an export file.
EXPORT_COLUMNS = [
    "id", "name", "description", "website", "domain", "email", "email_valid",
    "phone", "phone_normalized", "facebook", "linkedin", "twitter",
    "industry", "industries", "hq_location", "location_tail",
    "employee_enum", "employees_min", "employees_max",
    "funding_total_usd", "funding_rounds", "last_funding_type", "last_funding_at",
    "founded_on", "founded_year", "company_type", "operating_status",
    "ipo_status", "crunchbase_url", "quality_score", "lead_score",
    "verification_status", "source", "observed_at",
]
DEFAULT_EXPORT_COLUMNS = [
    "id", "name", "domain", "website", "industry", "hq_location",
    "employee_enum", "employees_min", "employees_max", "founded_year",
    "operating_status", "email", "phone", "linkedin",
    "quality_score", "lead_score", "verification_status",
]

EXPORT_MAX_ROWS = min(int(config.EXPORT_MAX_ROWS), 5000)  # contract cap


def _encode_cursor(sort_value: Any, row_id: str) -> str:
    raw = json.dumps({"s": sort_value, "i": row_id},
                     separators=(",", ":")).encode("utf-8")
    return base64.urlsafe_b64encode(raw).rstrip(b"=").decode("ascii")


def _decode_cursor(cursor: str, sort: str) -> tuple[Any, str]:
    try:
        padded = cursor + "=" * (-len(cursor) % 4)
        data = json.loads(base64.urlsafe_b64decode(padded.encode("ascii")).decode("utf-8"))
        sort_value, row_id = data["s"], data["i"]
    except Exception:
        raise APIError(422, "invalid_cursor", "cursor is malformed")
    if not isinstance(row_id, str) or not row_id:
        raise APIError(422, "invalid_cursor", "cursor is malformed")
    if sort == "name":
        if not isinstance(sort_value, str):
            raise APIError(422, "invalid_cursor", "cursor is malformed")
    else:
        try:
            sort_value = float(sort_value)
        except (TypeError, ValueError):
            raise APIError(422, "invalid_cursor", "cursor is malformed")
    return sort_value, row_id


def _validate_list_params(
    q: Optional[str],
    sort: str,
    dir_: Optional[str],
    limit: int,
    min_quality: Optional[float],
    employees_min: Optional[int],
    employees_max: Optional[int],
) -> str:
    if q is not None and len(q.strip()) < 2:
        raise APIError(422, "validation_error", "q must be at least 2 characters")
    if sort not in _SORT_EXPRS:
        raise APIError(422, "validation_error",
                       f"sort must be one of {sorted(_SORT_EXPRS)}")
    direction = (dir_ or _DEFAULT_DIRS.get(sort, "desc")).lower()
    if direction not in ("asc", "desc"):
        raise APIError(422, "validation_error", "dir must be 'asc' or 'desc'")
    if not 1 <= limit <= 100:
        raise APIError(422, "validation_error", "limit must be between 1 and 100")
    if min_quality is not None and not 0.0 <= min_quality <= 1.0:
        raise APIError(422, "validation_error", "min_quality must be between 0 and 1")
    if (employees_min is not None and employees_max is not None
            and employees_min > employees_max):
        raise APIError(422, "validation_error",
                       "employees_min must not exceed employees_max")
    return direction


def _company_filters(
    q: Optional[str] = None,
    industry: Optional[list[str]] = None,
    location: Optional[str] = None,
    employees_min: Optional[int] = None,
    employees_max: Optional[int] = None,
    has_email: Optional[bool] = None,
    has_phone: Optional[bool] = None,
    has_linkedin: Optional[bool] = None,
    min_quality: Optional[float] = None,
    operating_status: Optional[str] = None,
) -> tuple[list[str], list[Any]]:
    """Build parameterized WHERE clauses. Identifiers are static strings."""
    clauses: list[str] = []
    params: list[Any] = []

    if q:
        clauses.append("c.search_text ILIKE ?")
        params.append(f"%{q.strip()}%")
    if industry:
        vals = [v for v in industry if v]
        if vals:
            clauses.append(f"c.industry IN ({', '.join(['?'] * len(vals))})")
            params.extend(vals)
    if location:
        clauses.append("c.hq_location ILIKE ?")
        params.append(f"%{location.strip()}%")
    # Range overlap: NULL company bounds count as unbounded.
    if employees_min is not None:
        clauses.append("COALESCE(c.employees_max, 2147483647) >= ?")
        params.append(employees_min)
    if employees_max is not None:
        clauses.append("COALESCE(c.employees_min, 0) <= ?")
        params.append(employees_max)
    if has_email is True:
        clauses.append("c.email_valid IS TRUE")
    elif has_email is False:
        clauses.append("c.email_valid IS NOT TRUE")
    if has_phone is True:
        clauses.append("(c.phone_normalized IS NOT NULL AND c.phone_normalized <> '')")
    elif has_phone is False:
        clauses.append("NOT (c.phone_normalized IS NOT NULL AND c.phone_normalized <> '')")
    if has_linkedin is True:
        clauses.append("(c.linkedin IS NOT NULL AND c.linkedin <> '')")
    elif has_linkedin is False:
        clauses.append("NOT (c.linkedin IS NOT NULL AND c.linkedin <> '')")
    if min_quality is not None:
        clauses.append("c.quality_score >= ?")
        params.append(min_quality)
    if operating_status:
        clauses.append("c.operating_status = ?")
        params.append(operating_status.strip())

    return clauses, params


def _run_company_page(
    clauses: list[str],
    params: list[Any],
    sort: str,
    direction: str,
    limit: int,
    cursor: Optional[str],
) -> dict:
    sort_expr = _SORT_EXPRS[sort]
    dir_kw = "ASC" if direction == "asc" else "DESC"

    where = ""
    if clauses:
        where = "WHERE " + " AND ".join(clauses)

    page_params = list(params)
    if cursor:
        sort_value, last_id = _decode_cursor(cursor, sort)
        cmp_op = ">" if direction == "asc" else "<"
        keyset = (f"(({sort_expr}) {cmp_op} ? OR "
                  f"(({sort_expr}) = ? AND c.id > ?))")
        where = (where + " AND " + keyset) if where else "WHERE " + keyset
        page_params.extend([sort_value, sort_value, last_id])

    sql = (
        f"SELECT {SUMMARY_SELECT} FROM company c "
        f"{where} "
        f"ORDER BY {sort_expr} {dir_kw}, c.id ASC "
        f"LIMIT {limit + 1}"
    )
    with database.read_connection() as con:
        rows = con.execute(sql, page_params).fetchall()

    has_more = len(rows) > limit
    rows = rows[:limit]
    items = [dict(zip(_SUMMARY_KEYS, r)) for r in rows]
    next_cursor = None
    if has_more and rows:
        last = rows[-1]
        sort_col = {"name": "name", "quality_score": "quality_score",
                    "employees_max": "employees_max",
                    "founded_year": "founded_year"}[sort]
        next_cursor = _encode_cursor(last[_SUMMARY_KEYS.index(sort_col)], last[0])
    return {"items": items, "next_cursor": next_cursor, "limit": limit}


# ── Health ──────────────────────────────────────────────────────────────

@app.get("/health")
async def health():
    return {
        "status": "ok",
        "service": "leadengine-api",
        "version": API_VERSION,
        "db_rows": database.company_count(),
    }


# ── Auth ────────────────────────────────────────────────────────────────

@app.post("/api/v1/auth/signup", status_code=201)
async def signup(body: SignupIn):
    try:
        user = auth_service.signup(body.email, body.password, body.name)
    except ValueError as e:
        msg = str(e)
        if msg == "email already registered":
            raise APIError(409, "email_taken", "this email is already registered")
        raise APIError(422, "validation_error", msg)
    return {"user": user}


@app.post("/api/v1/auth/login")
async def login(body: LoginIn):
    try:
        result = auth_service.login(body.email, body.password)
    except ValueError:
        # Same message for unknown email and wrong password.
        raise APIError(401, "invalid_credentials", "invalid email or password")
    return result


@app.get("/api/v1/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return {"user": user}


# ── Companies ───────────────────────────────────────────────────────────

@app.get("/api/v1/companies", response_model=CompanyPage)
async def list_companies(
    q: Optional[str] = Query(default=None),
    industry: Optional[list[str]] = Query(default=None),
    location: Optional[str] = Query(default=None),
    employees_min: Optional[int] = Query(default=None, ge=0),
    employees_max: Optional[int] = Query(default=None, ge=0),
    has_email: Optional[bool] = Query(default=None),
    has_phone: Optional[bool] = Query(default=None),
    has_linkedin: Optional[bool] = Query(default=None),
    min_quality: Optional[float] = Query(default=None),
    operating_status: Optional[str] = Query(default=None),
    sort: str = Query(default="quality_score"),
    dir: Optional[str] = Query(default=None, alias="dir"),
    limit: int = Query(default=25),
    cursor: Optional[str] = Query(default=None),
):
    direction = _validate_list_params(q, sort, dir, limit, min_quality,
                                      employees_min, employees_max)
    clauses, params = _company_filters(
        q=q, industry=industry, location=location,
        employees_min=employees_min, employees_max=employees_max,
        has_email=has_email, has_phone=has_phone, has_linkedin=has_linkedin,
        min_quality=min_quality, operating_status=operating_status,
    )
    return _run_company_page(clauses, params, sort, direction, limit, cursor)


@app.get("/api/v1/companies/facets", response_model=FacetsOut)
async def company_facets():
    with database.read_connection() as con:
        industries = con.execute(
            "SELECT industry, COUNT(*) FROM company "
            "WHERE industry IS NOT NULL AND industry <> '' "
            "GROUP BY industry ORDER BY COUNT(*) DESC LIMIT 50"
        ).fetchall()
        locations = con.execute(
            "SELECT hq_location, COUNT(*) FROM company "
            "WHERE hq_location IS NOT NULL AND hq_location <> '' "
            "GROUP BY hq_location ORDER BY COUNT(*) DESC LIMIT 50"
        ).fetchall()
        statuses = con.execute(
            "SELECT operating_status, COUNT(*) FROM company "
            "WHERE operating_status IS NOT NULL AND operating_status <> '' "
            "GROUP BY operating_status ORDER BY COUNT(*) DESC"
        ).fetchall()
        ranges = con.execute(
            "SELECT employee_enum, COUNT(*) FROM company "
            "WHERE employee_enum IS NOT NULL AND employee_enum <> '' "
            "GROUP BY employee_enum ORDER BY COUNT(*) DESC"
        ).fetchall()
    return {
        "industries": [{"value": v, "count": c} for v, c in industries],
        "locations": [{"value": v, "count": c} for v, c in locations],
        "operating_statuses": [{"value": v, "count": c} for v, c in statuses],
        "employee_ranges": [
            {"value": v, "label": _humanize_enum(v), "count": c} for v, c in ranges
        ],
    }


def _humanize_enum(enum: str) -> str:
    """c_00001_00010 -> '1-10 employees'; unknown shapes pass through."""
    import re
    m = re.match(r"^c_0*(\d+)_0*(\d+)$", enum or "")
    if m:
        return f"{int(m.group(1)):,}-{int(m.group(2)):,} employees"
    return enum


def _export_rows(clauses: list[str], params: list[Any],
                 columns: list[str]) -> list[dict]:
    cols_sql = ", ".join(f'c."{c}"' for c in columns)  # c allowlisted by caller
    where = ("WHERE " + " AND ".join(clauses)) if clauses else ""
    sql = (f"SELECT {cols_sql} FROM company c {where} "
           f"ORDER BY c.id ASC LIMIT {EXPORT_MAX_ROWS + 1}")
    with database.read_connection() as con:
        rows = con.execute(sql, params).fetchall()
    if len(rows) > EXPORT_MAX_ROWS:
        raise APIError(400, "export_too_large",
                       f"export matches more than {EXPORT_MAX_ROWS} rows; "
                       "narrow the filters or export a smaller selection")
    return [dict(zip(columns, r)) for r in rows]


def _resolve_export_columns(columns: Optional[list[str]]) -> list[str]:
    if not columns:
        return DEFAULT_EXPORT_COLUMNS
    bad = [c for c in columns if c not in EXPORT_COLUMNS]
    if bad:
        raise APIError(422, "validation_error",
                       f"columns not allowlisted: {bad}")
    # Dedupe, keep requested order.
    return list(dict.fromkeys(columns))


def _build_export_file(rows: list[dict], columns: list[str],
                       fmt: str) -> tuple[bytes, str, str]:
    ts = int(time.time())
    if fmt == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(columns)
        for r in rows:
            writer.writerow(["" if r[c] is None else r[c] for c in columns])
        return (buf.getvalue().encode("utf-8"), "text/csv",
                f"leadengine-companies-{ts}.csv")
    if fmt == "json":
        payload = json.dumps({"exported_at": ts, "count": len(rows),
                              "items": rows}, default=str)
        return (payload.encode("utf-8"), "application/json",
                f"leadengine-companies-{ts}.json")
    # xlsx
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "companies"
    ws.append(columns)
    for r in rows:
        ws.append(["" if r[c] is None else r[c] for c in columns])
    buf = io.BytesIO()
    wb.save(buf)
    return (buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            f"leadengine-companies-{ts}.xlsx")


def _export_response(rows: list[dict], columns: list[str], fmt: str) -> StreamingResponse:
    body, media_type, filename = _build_export_file(rows, columns, fmt)
    return StreamingResponse(
        io.BytesIO(body),
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/v1/companies/export")
async def export_companies(
    q: Optional[str] = Query(default=None),
    industry: Optional[list[str]] = Query(default=None),
    location: Optional[str] = Query(default=None),
    employees_min: Optional[int] = Query(default=None, ge=0),
    employees_max: Optional[int] = Query(default=None, ge=0),
    has_email: Optional[bool] = Query(default=None),
    has_phone: Optional[bool] = Query(default=None),
    has_linkedin: Optional[bool] = Query(default=None),
    min_quality: Optional[float] = Query(default=None),
    operating_status: Optional[str] = Query(default=None),
    format: str = Query(default="csv", pattern="^(csv|json|xlsx)$"),
    columns: Optional[list[str]] = Query(default=None),
):
    if q is not None and len(q.strip()) < 2:
        raise APIError(422, "validation_error", "q must be at least 2 characters")
    if min_quality is not None and not 0.0 <= min_quality <= 1.0:
        raise APIError(422, "validation_error", "min_quality must be between 0 and 1")
    cols = _resolve_export_columns(columns)
    clauses, params = _company_filters(
        q=q, industry=industry, location=location,
        employees_min=employees_min, employees_max=employees_max,
        has_email=has_email, has_phone=has_phone, has_linkedin=has_linkedin,
        min_quality=min_quality, operating_status=operating_status,
    )
    rows = _export_rows(clauses, params, cols)
    return _export_response(rows, cols, format)


@app.post("/api/v1/companies/export")
async def export_companies_bulk(body: BulkExportIn):
    """Export an explicit list of company ids (bulk selection from the UI)."""
    cols = _resolve_export_columns(body.columns)
    placeholders = ", ".join(["?"] * len(body.ids))
    clauses = [f"c.id IN ({placeholders})"]
    rows = _export_rows(clauses, list(body.ids), cols)
    return _export_response(rows, cols, body.format)


@app.get("/api/v1/companies/{company_id}", response_model=CompanyDetail)
async def get_company(company_id: str):
    with database.read_connection() as con:
        row = con.execute("SELECT * FROM company WHERE id = ?", [company_id]).fetchone()
        if row is None:
            raise APIError(404, "not_found", "company not found")
        keys = [d[0] for d in con.description]
    rec = dict(zip(keys, row))

    industries: list[str] = []
    if rec.get("industries"):
        try:
            parsed = json.loads(rec["industries"])
            if isinstance(parsed, list):
                industries = [str(x) for x in parsed]
        except (json.JSONDecodeError, TypeError):
            industries = []
    signals: list[dict] = []
    if rec.get("score_signals"):
        try:
            parsed = json.loads(rec["score_signals"])
            if isinstance(parsed, list):
                signals = parsed
        except (json.JSONDecodeError, TypeError):
            signals = []

    return {
        "id": rec.get("id"),
        "name": rec.get("name"),
        "domain": rec.get("domain"),
        "website": rec.get("website"),
        "industry": rec.get("industry"),
        "hq_location": rec.get("hq_location"),
        "employee_enum": rec.get("employee_enum"),
        "employees_min": rec.get("employees_min"),
        "employees_max": rec.get("employees_max"),
        "founded_year": rec.get("founded_year"),
        "operating_status": rec.get("operating_status"),
        "has_email": rec.get("email_valid") is True,
        "has_phone": bool(rec.get("phone_normalized")),
        "has_linkedin": bool(rec.get("linkedin")),
        "quality_score": rec.get("quality_score"),
        "lead_score": rec.get("lead_score"),
        "verification_status": "unverified",
        "description": rec.get("description"),
        "email": rec.get("email"),
        "email_valid": rec.get("email_valid"),
        "phone": rec.get("phone"),
        "facebook": rec.get("facebook"),
        "linkedin": rec.get("linkedin"),
        "twitter": rec.get("twitter"),
        "industries": industries,
        "locations_raw": rec.get("locations_raw"),
        "location_tail": rec.get("location_tail"),
        "funding_total_usd": rec.get("funding_total_usd"),
        "funding_rounds": rec.get("funding_rounds"),
        "last_funding_type": rec.get("last_funding_type"),
        "last_funding_at": rec.get("last_funding_at"),
        "founded_on": rec.get("founded_on"),
        "company_type": rec.get("company_type"),
        "ipo_status": rec.get("ipo_status"),
        "crunchbase_url": rec.get("crunchbase_url"),
        "score_signals": signals,
        "source": rec.get("source"),
        "observed_at": rec.get("observed_at"),
    }


# ── Saved searches ──────────────────────────────────────────────────────

@app.get("/api/v1/saved-searches", response_model=SavedSearchList)
async def list_saved_searches():
    with database.read_connection() as con:
        rows = con.execute(
            "SELECT id, name, filters, created_at FROM saved_search "
            "ORDER BY created_at DESC"
        ).fetchall()
    items = []
    for sid, name, filters_json, created_at in rows:
        try:
            filters = json.loads(filters_json) if filters_json else {}
        except json.JSONDecodeError:
            filters = {}
        items.append({"id": sid, "name": name, "filters": filters,
                      "created_at": created_at})
    return {"items": items}


@app.post("/api/v1/saved-searches", status_code=201, response_model=SavedSearchOut)
async def create_saved_search(body: SavedSearchIn):
    name = body.name.strip()
    if not name:
        raise APIError(422, "validation_error", "name is required")
    if not isinstance(body.filters, dict):
        raise APIError(422, "validation_error", "filters must be an object")
    sid = str(uuid.uuid4())
    created_at = time.time()
    filters_json = json.dumps(body.filters)
    with database.write_connection() as con:
        con.execute(
            "INSERT INTO saved_search (id, name, filters, created_at) "
            "VALUES (?, ?, ?, ?)",
            [sid, name, filters_json, created_at],
        )
    return {"id": sid, "name": name, "filters": body.filters,
            "created_at": created_at}


@app.delete("/api/v1/saved-searches/{search_id}", status_code=204)
async def delete_saved_search(search_id: str):
    with database.write_connection() as con:
        exists = con.execute(
            "SELECT 1 FROM saved_search WHERE id = ?", [search_id]
        ).fetchone()
        if not exists:
            raise APIError(404, "not_found", "saved search not found")
        con.execute("DELETE FROM saved_search WHERE id = ?", [search_id])
    return None


# ── Favorites ───────────────────────────────────────────────────────────

@app.post("/api/v1/companies/{company_id}/favorites", status_code=201,
          response_model=FavoriteOut)
async def add_favorite(company_id: str):
    with database.read_connection() as con:
        exists = con.execute(
            "SELECT 1 FROM company WHERE id = ?", [company_id]
        ).fetchone()
    if not exists:
        raise APIError(404, "not_found", "company not found")
    with database.write_connection() as con:
        row = con.execute(
            "SELECT created_at FROM favorite WHERE company_id = ?", [company_id]
        ).fetchone()
        if row:
            created_at = row[0]
        else:
            created_at = time.time()
            con.execute(
                "INSERT INTO favorite (company_id, created_at) VALUES (?, ?)",
                [company_id, created_at],
            )
    return {"company_id": company_id, "created_at": created_at}


@app.delete("/api/v1/companies/{company_id}/favorites", status_code=204)
async def remove_favorite(company_id: str):
    with database.write_connection() as con:
        row = con.execute(
            "SELECT 1 FROM favorite WHERE company_id = ?", [company_id]
        ).fetchone()
        if not row:
            raise APIError(404, "not_found", "favorite not found")
        con.execute("DELETE FROM favorite WHERE company_id = ?", [company_id])
    return None


@app.get("/api/v1/favorites", response_model=CompanyPage)
async def list_favorites(
    limit: int = Query(default=25),
    cursor: Optional[str] = Query(default=None),
):
    if not 1 <= limit <= 100:
        raise APIError(422, "validation_error", "limit must be between 1 and 100")
    where = ""
    params: list[Any] = []
    if cursor:
        try:
            padded = cursor + "=" * (-len(cursor) % 4)
            data = json.loads(base64.urlsafe_b64decode(padded.encode("ascii"))
                              .decode("utf-8"))
            last_ts, last_id = float(data["s"]), str(data["i"])
        except Exception:
            raise APIError(422, "invalid_cursor", "cursor is malformed")
        where = ("WHERE (f.created_at < ? OR "
                 "(f.created_at = ? AND f.company_id > ?))")
        params = [last_ts, last_ts, last_id]
    sql = (
        f"SELECT {SUMMARY_SELECT} FROM favorite f "
        f"JOIN company c ON c.id = f.company_id "
        f"{where} "
        f"ORDER BY f.created_at DESC, f.company_id ASC "
        f"LIMIT {limit + 1}"
    )
    with database.read_connection() as con:
        rows = con.execute(sql, params).fetchall()
    has_more = len(rows) > limit
    rows = rows[:limit]
    items = [dict(zip(_SUMMARY_KEYS, r)) for r in rows]
    next_cursor = None
    if has_more and rows:
        # created_at is not in the summary projection; re-read it.
        with database.read_connection() as con:
            ts = con.execute(
                "SELECT created_at FROM favorite WHERE company_id = ?",
                [rows[-1][0]],
            ).fetchone()[0]
        next_cursor = _encode_cursor(ts, rows[-1][0])
    return {"items": items, "next_cursor": next_cursor, "limit": limit}
